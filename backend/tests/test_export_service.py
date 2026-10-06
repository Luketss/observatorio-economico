"""Exportacao XLSX generica (spec 2026-10-06-exportacao-graficos): slug, schema e planilha."""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.datas import agora_local
from app.core.slug import slugify
from app.schemas.export import ColunaExport, ExportXlsxIn
from app.services.export_service import MEDIA_XLSX, _celula, gerar_xlsx, nome_arquivo_xlsx


def _dados(**extra):
    base = dict(
        titulo="Evolução Anual do PIB",
        subtitulo="PIB total por ano",
        fonte="IBGE",
        municipio="Divinópolis",
        dataset="pib",
        colunas=[
            ColunaExport(chave="periodo", rotulo="Período", tipo="ano"),
            ColunaExport(chave="valor", rotulo="PIB Total", tipo="numero"),
            ColunaExport(chave="obs", rotulo="Obs", tipo="texto"),
        ],
        linhas=[
            {"periodo": 2021, "valor": 1234.5, "obs": "a; b"},
            {"periodo": "2022", "valor": "1500,25", "obs": None},
            {"periodo": 2023, "valor": None, "obs": "", "ignorada": 1},
        ],
    )
    base.update(extra)
    return ExportXlsxIn(**base)


# ---------- slugify ----------

def test_slugify_remove_acentos_e_pontuacao():
    assert slugify("Evolução Anual do PIB") == "evolucao-anual-do-pib"
    assert slugify("São João d'El Rei; MG") == "sao-joao-d-el-rei-mg"


def test_slugify_vazio_e_none():
    assert slugify("") == ""
    assert slugify(None) == ""
    assert slugify("---") == ""


def test_slugify_corta_no_max_len_sem_hifen_final():
    assert slugify("a" * 70, max_len=10) == "a" * 10
    assert slugify("abc def", max_len=4) == "abc"


# ---------- agora_local ----------

def test_agora_local_e_tz_aware_no_fuso_brasil():
    agora = agora_local()
    assert agora.tzinfo is not None
    assert agora.utcoffset().total_seconds() == -3 * 3600


# ---------- schema ----------

def test_schema_rejeita_mais_de_50_colunas():
    cols = [ColunaExport(chave=f"c{i}", rotulo=f"C{i}") for i in range(51)]
    with pytest.raises(ValidationError):
        ExportXlsxIn(titulo="t", colunas=cols, linhas=[])


def test_schema_rejeita_acima_de_50_mil_celulas():
    cols = [ColunaExport(chave=f"c{i}", rotulo=f"C{i}") for i in range(11)]
    linhas = [{"c0": 1}] * 4600  # 11 * 4600 = 50_600 > 50_000
    with pytest.raises(ValidationError):
        ExportXlsxIn(titulo="t", colunas=cols, linhas=linhas)


def test_schema_aceita_exatamente_o_limite():
    cols = [ColunaExport(chave=f"c{i}", rotulo=f"C{i}") for i in range(10)]
    dados = ExportXlsxIn(titulo="t", colunas=cols, linhas=[{"c0": 1}] * 5000)
    assert len(dados.linhas) == 5000


def test_schema_rejeita_tipo_desconhecido_e_sem_colunas():
    with pytest.raises(ValidationError):
        ColunaExport(chave="x", rotulo="X", tipo="data")
    with pytest.raises(ValidationError):
        ExportXlsxIn(titulo="t", colunas=[], linhas=[])


# ---------- _celula ----------

def test_celula_numero_aceita_string_com_virgula_e_devolve_none_para_vazio():
    assert _celula("1500,25", "numero") == 1500.25
    assert _celula(1234.5, "numero") == 1234.5
    assert _celula(None, "numero") is None
    assert _celula("", "numero") is None
    assert _celula("abc", "numero") == "abc"
    assert _celula(float("nan"), "numero") is None
    assert _celula(float("inf"), "numero") is None


def test_celula_ano_vira_int_e_texto_vira_str():
    assert _celula("2022", "ano") == 2022
    assert _celula(2023, "ano") == 2023
    assert _celula("n/d", "ano") == "n/d"
    assert _celula(42, "texto") == "42"


# ---------- gerar_xlsx ----------

def test_gerar_xlsx_estrutura_e_tipos():
    conteudo = gerar_xlsx(_dados())
    wb = load_workbook(BytesIO(conteudo))
    ws = wb["Dados"]
    assert ws["A1"].value == "Evolução Anual do PIB"
    assert ws["A1"].font.bold is True
    assert ws["A2"].value == "PIB total por ano"
    assert ws["A3"].value == "Município: Divinópolis"
    assert ws["A4"].value == "Fonte: IBGE"
    assert ws["A5"].value.startswith("Gerado em ") and ws["A5"].value.endswith("(UTC-3)")
    assert ws["A6"].value is None
    assert [c.value for c in ws[7]] == ["Período", "PIB Total", "Obs"]
    assert ws["A7"].font.bold is True
    # dados
    assert ws["A8"].value == 2021 and isinstance(ws["A8"].value, int)
    assert ws["B8"].value == 1234.5 and ws["B8"].number_format == "#,##0.00"
    assert ws["C8"].value == "a; b"
    assert ws["A9"].value == 2022  # "2022" virou int
    assert ws["B9"].value == 1500.25
    assert ws["C9"].value is None
    assert ws["B10"].value is None and ws["C10"].value is None
    # chave desconhecida ignorada: so 3 colunas
    assert ws.max_column == 3
    assert ws.auto_filter.ref == "A7:C10"


def test_gerar_xlsx_sem_metadados_opcionais_e_sem_linhas():
    dados = _dados(subtitulo=None, fonte=None, municipio=None, linhas=[])
    ws = load_workbook(BytesIO(gerar_xlsx(dados)))["Dados"]
    assert ws["A2"].value is None
    assert ws["A3"].value is None
    assert ws["A4"].value is None
    assert [c.value for c in ws[7]] == ["Período", "PIB Total", "Obs"]
    assert ws.auto_filter.ref == "A7:C7"


def test_gerar_xlsx_largura_de_coluna_limitada_a_60():
    longo = "x" * 200
    dados = _dados(linhas=[{"periodo": 2021, "valor": 1, "obs": longo}])
    ws = load_workbook(BytesIO(gerar_xlsx(dados)))["Dados"]
    assert ws.column_dimensions["C"].width == 60
    assert ws.column_dimensions["A"].width >= 8


def test_media_type_xlsx():
    assert MEDIA_XLSX == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------- nome_arquivo_xlsx ----------

def test_nome_arquivo_completo_e_sem_partes_vazias():
    nome = nome_arquivo_xlsx(_dados())
    assert nome.startswith("nid_pib_evolucao-anual-do-pib_divinopolis_")
    assert nome.endswith(".xlsx")
    assert len(nome.split("_")) == 5  # nid, dataset, painel, municipio, data
    nome2 = nome_arquivo_xlsx(_dados(dataset=None, municipio=None))
    assert nome2.startswith("nid_evolucao-anual-do-pib_")
    assert len(nome2.split("_")) == 3
