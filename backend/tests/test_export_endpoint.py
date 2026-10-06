"""POST /export/xlsx: handler, RBAC e Content-Disposition."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.api.deps import require_role
from app.api.v1.routers.export import exportar_xlsx
from app.core.exceptions import ForbiddenException
from app.schemas.export import ColunaExport, ExportXlsxIn
from app.services.export_service import MEDIA_XLSX


def _user(role):
    return SimpleNamespace(id=1, role=SimpleNamespace(nome=role), municipio_id=None)


def _dados():
    return ExportXlsxIn(
        titulo="Saldo CAGED",
        dataset="caged",
        municipio="Divinópolis",
        colunas=[
            ColunaExport(chave="periodo", rotulo="Período"),
            ColunaExport(chave="saldo", rotulo="Saldo", tipo="numero"),
        ],
        linhas=[{"periodo": "jan/26", "saldo": 120}, {"periodo": "fev/26", "saldo": -30}],
    )


def test_handler_devolve_xlsx_com_content_disposition():
    resp = exportar_xlsx(_dados(), current_user=_user("ADMIN_GLOBAL"))
    assert resp.status_code == 200
    assert resp.media_type == MEDIA_XLSX
    cd = resp.headers["content-disposition"]
    assert cd.startswith('attachment; filename="nid_caged_saldo-caged_divinopolis_')
    assert cd.endswith('.xlsx"')
    ws = load_workbook(BytesIO(resp.body))["Dados"]
    assert ws["A1"].value == "Saldo CAGED"
    assert ws["B9"].value == -30


def test_rbac_admin_municipio_e_visualizador_sao_barrados():
    dep = require_role("ADMIN_GLOBAL")
    for papel in ("ADMIN_MUNICIPIO", "VISUALIZADOR", "ANALISTA"):
        with pytest.raises(ForbiddenException):
            dep(current_user=_user(papel))


def test_rbac_admin_global_passa():
    dep = require_role("ADMIN_GLOBAL")
    user = _user("ADMIN_GLOBAL")
    assert dep(current_user=user) is user


def test_rota_registrada_no_app_com_prefixo_e_dependencia_de_papel():
    from app.main import app

    rotas = {r.path: r for r in app.routes if hasattr(r, "path")}
    assert "/api/v1/export/xlsx" in rotas
    rota = rotas["/api/v1/export/xlsx"]
    assert rota.methods == {"POST"}
    nomes = [getattr(d.call, "__qualname__", "") for d in rota.dependant.dependencies]
    assert any("role_checker" in n for n in nomes), nomes
