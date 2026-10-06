"""Gera a planilha XLSX generica do menu Exportar (so ADMIN_GLOBAL, ver router export).

Nao sabe o que e PIB ou CAGED: tabula o que recebe. Aba "Dados": titulo (A1),
subtitulo (A2), municipio (A3), fonte (A4), gerado em (A5), linha em branco,
cabecalho na linha 7 com filtro automatico, dados a partir da linha 8.
"""
import math
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.datas import agora_local
from app.core.slug import slugify
from app.schemas.export import ExportXlsxIn

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LINHA_CABECALHO = 7
LARGURA_MIN = 8
LARGURA_MAX = 60


def _escrever_texto(celula, texto: str) -> None:
    """Escreve texto e força data_type='s' para evitar injecao de formulas (=...).

    Tambem remove caracteres XML ilegais.
    """
    texto_limpo = ILLEGAL_CHARACTERS_RE.sub("", texto)
    celula.value = texto_limpo
    celula.data_type = "s"


def _celula(valor, tipo: str):
    """Converte o valor cru do front para o tipo da coluna. Vazio/nao-finito -> None."""
    if valor is None or valor == "":
        return None
    if tipo == "numero":
        if isinstance(valor, bool):
            return int(valor)
        if isinstance(valor, (int, float)):
            if isinstance(valor, float) and not math.isfinite(valor):
                return None
            return valor
        try:
            valor_str = str(valor)
            if "," in valor_str:
                valor_str = valor_str.replace(".", "").replace(",", ".")
            parsed = float(valor_str)
            if not math.isfinite(parsed):
                return None
            return parsed
        except ValueError:
            return str(valor)
    if tipo == "ano":
        try:
            return int(valor)
        except (TypeError, ValueError, OverflowError):
            return str(valor)
    return str(valor)


def nome_arquivo_xlsx(dados: ExportXlsxIn) -> str:
    partes = [
        "nid",
        slugify(dados.dataset),
        slugify(dados.titulo),
        slugify(dados.municipio),
        agora_local().strftime("%Y-%m-%d"),
    ]
    return "_".join(p for p in partes if p) + ".xlsx"


def gerar_xlsx(dados: ExportXlsxIn) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"

    _escrever_texto(ws["A1"], dados.titulo)
    ws["A1"].font = Font(bold=True, size=14)
    if dados.subtitulo:
        _escrever_texto(ws["A2"], dados.subtitulo)
    if dados.municipio:
        _escrever_texto(ws["A3"], f"Município: {dados.municipio}")
    if dados.fonte:
        _escrever_texto(ws["A4"], f"Fonte: {dados.fonte}")
    _escrever_texto(ws["A5"], "Gerado em " + agora_local().strftime("%d/%m/%Y %H:%M") + " (UTC-3)")

    negrito = Font(bold=True)
    cinza = PatternFill("solid", fgColor="EEEEEE")
    larguras = []
    for ci, col in enumerate(dados.colunas, start=1):
        celula = ws.cell(row=LINHA_CABECALHO, column=ci)
        _escrever_texto(celula, col.rotulo)
        celula.font = negrito
        celula.fill = cinza
        larguras.append(len(col.rotulo))

    for ri, linha in enumerate(dados.linhas, start=LINHA_CABECALHO + 1):
        for ci, col in enumerate(dados.colunas, start=1):
            valor = _celula(linha.get(col.chave), col.tipo)
            celula = ws.cell(row=ri, column=ci)
            if isinstance(valor, str):
                _escrever_texto(celula, valor)
            else:
                celula.value = valor
            if col.tipo == "numero" and isinstance(valor, (int, float)):
                celula.number_format = "#,##0.00"
            elif col.tipo == "ano" and isinstance(valor, int):
                celula.number_format = "0"
            if valor is not None:
                larguras[ci - 1] = max(larguras[ci - 1], len(str(valor)))

    for ci, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = min(LARGURA_MAX, max(LARGURA_MIN, largura + 2))

    ultima_linha = LINHA_CABECALHO + len(dados.linhas)
    ws.auto_filter.ref = f"A{LINHA_CABECALHO}:{get_column_letter(len(dados.colunas))}{ultima_linha}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
